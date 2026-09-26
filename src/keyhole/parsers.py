"""Isolated document parser subprocess: one validated anonymous snapshot in, JSON out.

Runs as ``python -I -m keyhole.parsers`` under CPU and output budgets, with a
best-effort address-space limit. Process separation is not an OS sandbox.
No URLs, macros or project code are intentionally executed.
"""

import argparse
import base64
import contextlib
import io
import json
import math
import os
import resource
import warnings
import zipfile

TEXT_BUDGET = 60 * 1024
PIXEL_BUDGET = 25_000_000
OUTPUT_PIXELS = 2_000_000
OUTPUT_IMAGE_BYTES = 3 * 1024 * 1024


class Refused(Exception):
    def __init__(self, code, message):
        self.code, self.message = code, message


def require(value, code, message):
    if not value:
        raise Refused(code, message)


def zip_check(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        require(len(entries) <= 20_000, "archive_budget", "Document has too many archive entries.")
        total = 0
        for item in entries:
            total += item.file_size
            require(
                total <= 256 * 1024 * 1024 and item.file_size <= 64 * 1024 * 1024,
                "archive_budget",
                "Expanded document exceeds its byte budget.",
            )
            require(
                not item.flag_bits & 1, "encrypted_document", "Encrypted documents are unsupported."
            )
            require(
                item.file_size <= max(1024 * 1024, item.compress_size * 200),
                "archive_budget",
                "Document compression ratio exceeds its budget.",
            )
            require(
                not item.filename.startswith("/") and ".." not in item.filename.split("/"),
                "invalid_document",
                "Unsafe archive member path.",
            )
            if item.filename.lower().endswith((".xml", ".rels")):
                xml = archive.read(item)
                require(
                    b"\x00" not in xml,
                    "invalid_document",
                    "OOXML must use UTF-8 XML without external entities.",
                )
                lowered = xml.lower()
                require(
                    b"<!doctype" not in lowered and b"<!entity" not in lowered,
                    "invalid_document",
                    "XML entities and DTDs are unsupported.",
                )


def bounded_units(units, unit, start, limit, total=None, source_start=1):
    picked, used, partial = [], 0, False
    count = 0
    for number, value in enumerate(units, source_start):
        count = number
        if number < start:
            continue
        if number >= start + limit:
            break
        text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
        encoded = text.encode()
        if used + len(encoded) > TEXT_BUDGET:
            if not picked:
                picked.append(
                    {
                        unit: number,
                        "text": encoded[: TEXT_BUDGET - 1024].decode("utf-8", "ignore"),
                        "partial_unit": True,
                    }
                )
                partial = True
            break
        picked.append({unit: number, "text": text})
        used += len(encoded)
    total = total if total is not None else count
    require(
        start <= max(1, total), "invalid_range", "Requested starting unit is outside the document."
    )
    last = picked[-1][unit] if picked else 0
    return {
        "range": {"unit": unit, "start": start, "end": last, "total": total},
        "items": picked,
        "truncated": partial or last < total,
        "partial_unit": partial,
        "next_start": last + 1 if last < total else None,
        "note": "Oversized unit was partially extracted." if partial else None,
    }


def image_result(image):
    from PIL import Image, ImageOps

    require(
        image.width * image.height <= PIXEL_BUDGET,
        "image_budget",
        "Image dimensions exceed the pixel budget.",
    )
    original = [image.width, image.height]
    image = ImageOps.exif_transpose(image)
    factor = min(1, math.sqrt(OUTPUT_PIXELS / max(1, image.width * image.height)))
    image = (
        image.resize((max(1, int(image.width * factor)), max(1, int(image.height * factor))))
        if factor < 1
        else image.copy()
    )
    if image.mode not in ("RGB", "RGBA", "L", "LA"):
        image = image.convert("RGBA")
    stream = io.BytesIO()
    image.save(stream, "PNG")
    mime = "image/png"
    if len(stream.getvalue()) > OUTPUT_IMAGE_BYTES:
        background = Image.new("RGB", image.size, "white")
        if "A" in image.getbands():
            background.paste(image.convert("RGB"), mask=image.getchannel("A"))
        else:
            background.paste(image.convert("RGB"))
        stream = io.BytesIO()
        background.save(stream, "JPEG", quality=85)
        mime = "image/jpeg"
    require(
        len(stream.getvalue()) <= OUTPUT_IMAGE_BYTES,
        "image_budget",
        "Rendered image exceeds its response budget.",
    )
    return {
        "image": {"data": base64.b64encode(stream.getvalue()).decode("ascii"), "mime_type": mime},
        "original_dimensions": original,
        "returned_dimensions": [image.width, image.height],
        "truncated": False,
        "transformed": True,
        "note": "Display image is normalized and may be resized; source SHA-256 refers to the original bytes. Metadata is not copied.",
    }


def parse(data, kind, options):
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = PIXEL_BUDGET
    warnings.simplefilter("error", Image.DecompressionBombWarning)
    mode, start, limit = (
        options.get("mode", "auto"),
        options.get("start", 1),
        options.get("limit", 0),
    )
    require(start >= 1 and limit >= 0, "invalid_range", "Range values must be positive.")
    if kind == "image":
        with Image.open(io.BytesIO(data)) as img:
            require(
                img.format in ("PNG", "JPEG", "WEBP"),
                "unsupported_image",
                "Only PNG, JPEG and WebP are supported.",
            )
            result = image_result(img)
            result.update(
                {"format": "image", "frame": 1, "frame_count": getattr(img, "n_frames", 1)}
            )
            return result
    if kind in ("docx", "pptx", "xlsx"):
        zip_check(data)
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data), strict=True)
        require(not reader.is_encrypted, "encrypted_document", "Encrypted PDFs are unsupported.")
        total = len(reader.pages)
        require(
            total <= 5000 and 1 <= start <= max(1, total),
            "document_budget",
            "Page count or starting page exceeds the budget.",
        )
        if mode == "image":
            import pypdfium2 as pdfium

            with pdfium.PdfDocument(data) as pdf:
                page = pdf[start - 1]
                w, h = page.get_size()
                require(
                    w > 0 and h > 0 and w * h < 1e10, "image_budget", "Invalid PDF page dimensions."
                )
                scale = min(2, math.sqrt(OUTPUT_PIXELS / (w * h)))
                bitmap = page.render(scale=scale, may_draw_forms=False)
                try:
                    result = image_result(bitmap.to_pil())
                finally:
                    bitmap.close()
                    page.close()
                result.update(
                    {
                        "format": "pdf",
                        "range": {"unit": "page", "start": start, "end": start, "total": total},
                    }
                )
                return result
        limit = limit or 1
        require(limit <= 5, "invalid_range", "Read at most 5 PDF pages per call.")
        pages = (
            reader.pages[i].extract_text() or ""
            for i in range(start - 1, min(total, start - 1 + limit))
        )
        result = bounded_units(pages, "page", start, limit, total, source_start=start)
        result.update(
            {
                "format": "pdf",
                "ocr_performed": False,
                "scan_note": "Empty page text may require mode=image; no OCR is performed.",
            }
        )
        return result
    if kind == "docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(io.BytesIO(data))
        units = []
        for item in document.element.body.iterchildren():
            if item.tag.endswith("}p"):
                units.append(Paragraph(item, document).text)
            elif item.tag.endswith("}tbl"):
                units.append(
                    "\n".join(
                        "\t".join(cell.text for cell in row.cells)
                        for row in Table(item, document).rows
                    )
                )
        require(
            (limit or 100) <= 400, "invalid_range", "Read at most 400 document blocks per call."
        )
        result = bounded_units(units, "block", start, limit or 100, len(units))
        result.update(
            {
                "format": "docx",
                "extraction": "Body paragraphs and tables in document order; drawings, comments, footnotes, headers and layout are not extracted.",
            }
        )
        return result
    if kind == "pptx":
        from pptx import Presentation

        presentation = Presentation(io.BytesIO(data))
        require((limit or 5) <= 20, "invalid_range", "Read at most 20 slides per call.")

        def shapes_text(shapes):
            from pptx.enum.shapes import MSO_SHAPE_TYPE

            for shape in shapes:
                if shape.has_text_frame:
                    yield shape.text
                if shape.has_table:
                    yield "\n".join(
                        "\t".join(c.text for c in row.cells) for row in shape.table.rows
                    )
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    yield from shapes_text(shape.shapes)

        total = len(presentation.slides)
        units = (
            "\n".join(shapes_text(presentation.slides[i].shapes))
            for i in range(start - 1, min(total, start - 1 + (limit or 5)))
        )
        result = bounded_units(units, "slide", start, limit or 5, total, source_start=start)
        result.update(
            {
                "format": "pptx",
                "extraction": "Slide and table text, including grouped shapes; speaker notes, charts, images and layout are not extracted.",
            }
        )
        return result
    if kind == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.cell import get_column_letter, range_boundaries

        formulas = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        values = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            sheet = options.get("sheet") or formulas.sheetnames[0]
            require(
                sheet in formulas.sheetnames, "invalid_sheet", "Requested sheet does not exist."
            )
            selected = options.get("cell_range") or "A1:T50"
            left, top, right, bottom = range_boundaries(selected)
            require(
                all(isinstance(v, int) for v in (left, top, right, bottom))
                and 1 <= left <= right <= 16384
                and 1 <= top <= bottom <= 1_048_576
                and (right - left + 1) * (bottom - top + 1) <= 2000,
                "invalid_range",
                "Use a bounded A1 range containing at most 2000 cells.",
            )
            rows, budget, partial, last_row = [], 0, False, top - 1
            frows = formulas[sheet].iter_rows(
                min_row=top, max_row=bottom, min_col=left, max_col=right
            )
            vrows = values[sheet].iter_rows(
                min_row=top, max_row=bottom, min_col=left, max_col=right
            )
            for number, (frow, vrow) in enumerate(zip(frows, vrows, strict=False), top):
                row = []
                for col, (fcell, vcell) in enumerate(zip(frow, vrow, strict=False), left):
                    value = fcell.value
                    formula = value if fcell.data_type == "f" else None
                    row.append(
                        {
                            "cell": f"{get_column_letter(col)}{number}",
                            "formula": formula,
                            "value": vcell.value if formula is not None else value,
                            "cached_value_available": vcell.value is not None
                            if formula is not None
                            else None,
                        }
                    )
                encoded = json.dumps(row, ensure_ascii=False, default=str).encode()
                if budget + len(encoded) > TEXT_BUDGET:
                    partial = True
                    break
                rows.append(row)
                last_row = number
                budget += len(encoded)
            require(
                rows or not partial,
                "unit_too_large",
                "One row exceeds the response budget; request fewer columns.",
            )
            return {
                "format": "xlsx",
                "sheets": formulas.sheetnames[:200],
                "sheet_count": len(formulas.sheetnames),
                "sheet": sheet,
                "range": {
                    "requested": selected,
                    "start_row": top,
                    "end_row": last_row,
                    "start_column": left,
                    "end_column": right,
                },
                "rows": rows,
                "truncated": partial,
                "next_range": f"{get_column_letter(left)}{last_row + 1}:{get_column_letter(right)}{bottom}"
                if partial
                else None,
                "recalculated": False,
                "external_links_loaded": False,
                "note": "Formula values are existing caches, possibly stale or missing. No formulas or macros are executed.",
            }
        finally:
            formulas.close()
            values.close()
    raise Refused("unsupported_format", "Only metadata is available for this format.")


def main() -> None:
    parser = argparse.ArgumentParser(prog="keyhole.parsers")
    parser.add_argument("--input-fd", type=int, required=True)
    parser.add_argument("--kind", required=True)
    parser.add_argument("--generation", required=True)
    parser.add_argument("--nonce", required=True)
    parser.add_argument("--state-dir", help="owning runtime identity; never opened by the parser")
    args = parser.parse_args()
    resource.setrlimit(resource.RLIMIT_CPU, (15, 16))
    resource.setrlimit(resource.RLIMIT_CORE, (0, 0))
    # macOS can reject RLIMIT_AS. CPU, wall-clock and input limits still apply;
    # this process has the user's OS permissions and is not a network/filesystem sandbox.
    with contextlib.suppress(OSError, ValueError):
        resource.setrlimit(resource.RLIMIT_AS, (2 * 1024**3, 2 * 1024**3))
    try:
        with os.fdopen(args.input_fd, "rb") as source:
            data = source.read(64 * 1024 * 1024 + 1)
        require(
            len(data) <= 64 * 1024 * 1024,
            "file_too_large",
            "Source exceeds the parser input budget.",
        )
        options = json.loads(os.read(0, 16384))
        result = parse(data, args.kind, options)
    except Refused as exc:
        result = {"error": {"code": exc.code, "message": exc.message}}
    except Exception:
        result = {
            "error": {
                "code": "invalid_document",
                "message": "This document is damaged, unsupported, or exceeds parser limits; content was not inferred.",
            }
        }
    print(json.dumps(result, ensure_ascii=False, default=str))


if __name__ == "__main__":
    main()
