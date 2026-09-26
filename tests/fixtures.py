"""Synthetic document, image and text fixtures for the reader tests."""

import zipfile
from pathlib import Path

MARKER_PREFIX = "KEYHOLE-SYNTHETIC-"
UNICODE_LINE = "Ünïcödé prüfung: Строка два ✓"


def make(root: Path, label: str) -> dict:
    from docx import Document
    from openpyxl import Workbook
    from PIL import Image, ImageDraw
    from pptx import Presentation
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    root.mkdir(parents=True, exist_ok=True)
    marker = MARKER_PREFIX + label
    (root / "summary.md").write_text(
        f"# {label}\n\nSynthetic fixtures only.\nMarker: {marker}\n{UNICODE_LINE}\nState: version-one\n",
        encoding="utf-8",
    )
    nested = root / "nested"
    nested.mkdir(exist_ok=True)
    (nested / "notes.txt").write_text(
        f"{label} nested evidence\nLiteral needle: bridge-verification\n"
    )

    doc = Document()
    doc.add_paragraph(f"{label} document evidence")
    doc.add_paragraph(UNICODE_LINE)
    table = doc.add_table(rows=2, cols=2)
    for row, texts in zip(table.rows, (("Item", "Value"), ("Marker", marker)), strict=True):
        for cell, text in zip(row.cells, texts, strict=True):
            cell.text = text
    doc.save(root / "sample.docx")

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = f"{label} slide"
    slide.placeholders[1].text = UNICODE_LINE + "\n" + marker
    deck.save(root / "sample.pptx")

    book = Workbook()
    sheet = book.active
    sheet.title = "Evidence"
    sheet.append(["Label", "Input A", "Input B", "Total"])
    sheet.append([label, 7, 11, "=B2+C2"])
    sheet.append([UNICODE_LINE, marker])
    book.save(root / "sample.xlsx")
    # Inject a cached formula value the way a spreadsheet application would have saved it.
    with zipfile.ZipFile(root / "sample.xlsx") as original:
        entries = {i.filename: original.read(i) for i in original.infolist()}
    entries["xl/worksheets/sheet1.xml"] = entries["xl/worksheets/sheet1.xml"].replace(
        b"<f>B2+C2</f><v></v>", b"<f>B2+C2</f><v>18</v>"
    )
    with zipfile.ZipFile(root / "sample.xlsx", "w", zipfile.ZIP_DEFLATED) as out:
        for name, value in entries.items():
            out.writestr(name, value)

    image = Image.new("RGB", (900, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((50, 80, 280, 360), fill="#eb4934")
    draw.ellipse((370, 100, 630, 360), fill="#2a70d6")
    draw.polygon([(750, 80), (650, 360), (860, 360)], fill="#26a66b")
    draw.text((50, 30), f"{label}: red rectangle / blue circle / green triangle", fill="black")
    for ext in ("png", "jpg", "webp"):
        image.save(root / f"shapes.{ext}")
    image.save(root / "scanned.pdf", "PDF", resolution=100)

    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        f"BT /F1 18 Tf 50 720 Td ({label} PDF evidence) Tj 0 -30 Td ({marker}) Tj ET".encode()
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(root / "sample.pdf")
    (root / "damaged.pdf").write_bytes(b"%PDF-1.7\nnot a valid document")
    (root / "unknown.bin").write_bytes(b"\x00\xff\x89binary")
    return {"path": str(root), "label": label, "marker": marker, "synthetic_only": True}
